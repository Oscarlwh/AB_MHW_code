from __future__ import annotations

import csv
from pathlib import Path

from docx import Document


SOURCE_DOCX = Path(
    "/path/to/user/Documents/DesktopOrganizer/Folders/Paper/论文/稿子/"
    "V3_SI_AB_MHW.docx"
)
SOURCE_CSV = Path(
    "/path/to/user/Documents/DesktopOrganizer/Folders/Paper/论文/表格文件/"
    "NEP_MHW_events_1981-2024_area_mean_smoothed30.csv"
)
OUTPUT_DOCX = SOURCE_DOCX.with_name("V3_SI_AB_MHW_TableS1_updated.docx")

PEAK_DATES = [
    "1985/11/29",
    "1986/11/19",
    "1989/11/18",
    "1989/12/24",
    "1991/11/10",
    "1993/11/26",
    "2004/12/24",
    "2015/11/6",
    "2015/11/30",
    "2018/11/19",
    "2019/11/10",
    "2020/11/13",
    "2023/11/3",
    "2023/12/14",
]


def set_cell_text_preserving_format(cell, value: str) -> None:
    """Replace visible cell text while retaining its paragraph/run styling."""
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def load_selected_rows() -> list[list[str]]:
    with SOURCE_CSV.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    by_peak = {row["peak_date"]: row for row in rows}
    missing = [peak for peak in PEAK_DATES if peak not in by_peak]
    if missing:
        raise ValueError(f"Missing selected peak dates in CSV: {missing}")

    selected = []
    for peak in PEAK_DATES:
        row = by_peak[peak]
        selected.append(
            [
                row["start_date"],
                row["end_date"],
                row["peak_date"],
                str(int(float(row["duration_days"]))),
                f"{float(row['max_intensity']):.2f}",
                f"{float(row['cum_intensity']):.2f}",
            ]
        )
    return selected


def find_table_s1(document: Document):
    for table in document.tables:
        if len(table.rows) != 15 or len(table.columns) != 6:
            continue
        header = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
        if header[0] == "Start date" and header[2] == "Peak date":
            return table
    raise ValueError("Could not locate the 15-row, 6-column Table S1")


def main() -> None:
    selected = load_selected_rows()
    document = Document(SOURCE_DOCX)
    table = find_table_s1(document)

    for row_index, values in enumerate(selected, start=1):
        for column_index, value in enumerate(values):
            set_cell_text_preserving_format(
                table.rows[row_index].cells[column_index],
                value,
            )

    document.save(OUTPUT_DOCX)

    # Structural verification after the round trip.
    check = Document(OUTPUT_DOCX)
    check_table = find_table_s1(check)
    actual = [
        [cell.text.strip() for cell in row.cells]
        for row in check_table.rows[1:]
    ]
    if actual != selected:
        raise RuntimeError("Table S1 verification failed after saving")

    print(OUTPUT_DOCX)
    for values in selected:
        print(" | ".join(values))


if __name__ == "__main__":
    main()
